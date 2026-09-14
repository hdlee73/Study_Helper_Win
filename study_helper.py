"""Study Helper 2.2 application entry point.

Keeps the proven audio/TTS/Anki backend in :mod:`study_helper_core` while adding
rate-limited translation, DOCX transcript export with optional timelines, and a
refreshed desktop UI.
"""
from __future__ import annotations

import logging
import multiprocessing
import os
import queue
import sys
import threading
import time
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Tuple

import customtkinter as ctk
from tkinter import messagebox

import study_helper_core as core

# Re-export backend helpers used by smoke/regression tests and existing imports.
APP_DIR = core.APP_DIR
CONFIG_PATH = core.CONFIG_PATH
RESOURCE_DIR = core.RESOURCE_DIR
atomic_write = core.atomic_write
dedupe_df = core.dedupe_df
detect_lang = core.detect_lang
get_whisper_model = core.get_whisper_model
pandas_module = core.pandas_module
sanitize_filename = core.sanitize_filename
save_rows_to_excel = core.save_rows_to_excel
split_sentences = core.split_sentences

APP_VERSION = "2.2.0-rc.1"

PALETTE = {
    "page": "#F4F7FB",
    "sidebar": "#0F172A",
    "sidebar_hover": "#1E293B",
    "sidebar_text": "#CBD5E1",
    "sidebar_muted": "#94A3B8",
    "primary": "#2563EB",
    "primary_hover": "#1D4ED8",
    "primary_soft": "#EFF6FF",
    "text": "#0F172A",
    "muted": "#64748B",
    "border": "#E2E8F0",
    "input": "#F8FAFC",
    "card": "#FFFFFF",
    "success_bg": "#ECFDF5",
    "success": "#047857",
    "error_bg": "#FEF2F2",
    "error": "#B91C1C",
}

# ---------------------------------------------------------------------------
# Speech-to-text with timestamps
# ---------------------------------------------------------------------------
def transcribe_audio(path: str):
    """Transcribe audio while retaining Whisper segment start/end timestamps."""
    segments, _info = core.get_whisper_model().transcribe(
        str(Path(path).resolve()),
        beam_size=max(1, int(core.CONFIG.get("whisper_beam_size", 1))),
        vad_filter=True,
        condition_on_previous_text=False,
    )
    items = []
    for segment in segments:
        text = (segment.text or "").strip()
        if not text:
            continue
        start = max(0.0, float(getattr(segment, "start", 0.0) or 0.0))
        end = max(start, float(getattr(segment, "end", start) or start))
        items.append({"text": text, "start": start, "end": end})
    if not items:
        raise RuntimeError("음성 인식 결과가 비어 있습니다.")
    return {"text": " ".join(item["text"] for item in items), "segments": items}


# Make legacy backend callers (raw Excel export etc.) benefit from timestamps too.
core.transcribe_audio = transcribe_audio

# ---------------------------------------------------------------------------
# Translation: throttle + retry to prevent Google 5 requests/second failures
# ---------------------------------------------------------------------------
_TRANSLATION_LOCK = threading.Lock()
_LAST_TRANSLATION_REQUEST = 0.0
_TRANSLATORS: Dict[Tuple[str, str], object] = {}


def _translation_min_interval() -> float:
    # 0.30 sec => max ~3.3 req/s, leaving headroom under the reported 5 req/s cap.
    try:
        return max(0.25, float(core.CONFIG.get("translation_min_interval", 0.30)))
    except Exception:
        return 0.30


def _translator(source: str, target: str):
    key = (source, target)
    if key not in _TRANSLATORS:
        from deep_translator import GoogleTranslator
        _TRANSLATORS[key] = GoogleTranslator(source=source, target=target)
    return _TRANSLATORS[key]


def _throttled_translate(text: str, source: str, target: str) -> str:
    """Translate one sentence with global pacing and exponential backoff."""
    global _LAST_TRANSLATION_REQUEST
    last_error = None
    for attempt in range(5):
        try:
            # A single lock protects both pacing and the translator instance from
            # concurrent calls started by older code paths.
            with _TRANSLATION_LOCK:
                wait = _translation_min_interval() - (time.monotonic() - _LAST_TRANSLATION_REQUEST)
                if wait > 0:
                    time.sleep(wait)
                _LAST_TRANSLATION_REQUEST = time.monotonic()
                result = _translator(source, target).translate(text)
            if result is None or not str(result).strip():
                raise RuntimeError("번역 결과가 비어 있습니다.")
            return str(result).strip()
        except Exception as exc:
            last_error = exc
            message = str(exc).lower()
            rate_limited = any(token in message for token in (
                "too many requests", "429", "server error", "request limit", "rate limit"
            ))
            # Rate-limit errors need a longer cool-down. Transient network errors
            # also get a short retry, but programming/data errors are not hidden.
            if attempt == 4:
                break
            if rate_limited:
                time.sleep(min(16.0, 2.0 ** (attempt + 1)))
            elif attempt < 2:
                time.sleep(1.0 + attempt)
            else:
                raise
    raise RuntimeError(
        "Google 번역 요청 제한에 걸렸습니다. 자동 속도 조절과 재시도를 했지만 번역하지 못했습니다. "
        "잠시 후 다시 시도해 주세요."
    ) from last_error


@lru_cache(maxsize=4096)
def translate_pair(sentence: str, lang: str) -> Tuple[str, str]:
    if lang == "ko":
        return sentence, _throttled_translate(sentence, "ko", "en")
    return _throttled_translate(sentence, "en", "ko"), sentence


def transcribe_mp3_to_rows(mp3_path: str):
    """Transcribe then translate sequentially to stay safely below free endpoint limits."""
    result = transcribe_audio(mp3_path)
    texts = [
        seg.get("text", "") for seg in result.get("segments", [])
        if isinstance(seg, dict) and seg.get("text", "").strip()
    ]
    if not texts:
        texts = [result.get("text", "")]
    sentences = split_sentences(texts)
    if not sentences:
        raise RuntimeError("문장으로 분리할 수 있는 음성 인식 결과가 없습니다.")

    stem = Path(mp3_path).stem
    rows = []
    for sentence in sentences:
        front, back = translate_pair(sentence, detect_lang(sentence))
        rows.append([front, back, stem])
    return rows


core.translate_pair = translate_pair
core.transcribe_mp3_to_rows = transcribe_mp3_to_rows

# ---------------------------------------------------------------------------
# DOCX transcript generation
# ---------------------------------------------------------------------------
def _format_time(seconds: float) -> str:
    total = max(0, int(round(float(seconds))))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _set_east_asia_font(run, name: str = "Malgun Gothic"):
    from docx.oxml.ns import qn
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def speech_docx(segments, target, include_timeline: bool = True, title: str = ""):
    """Write a clean Word transcript.

    With timelines enabled, timestamps live only in the first table column. In
    Word the user can select that column and choose Delete Columns to remove all
    timestamps at once without touching transcript text.
    """
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    target = Path(target)
    normalized = []
    for item in segments or []:
        if isinstance(item, dict):
            text = str(item.get("text", "")).strip()
            start = float(item.get("start", 0.0) or 0.0)
            end = float(item.get("end", start) or start)
        else:
            text = str(item).strip()
            start = end = 0.0
        if text:
            normalized.append({"text": text, "start": start, "end": max(start, end)})
    if not normalized:
        raise RuntimeError("DOCX로 저장할 전사 내용이 없습니다.")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    if "Timeline" not in [style.name for style in doc.styles]:
        timeline_style = doc.styles.add_style("Timeline", WD_STYLE_TYPE.PARAGRAPH)
        timeline_style.font.name = "Malgun Gothic"
        timeline_style.font.size = Pt(9)
        timeline_style.font.color.rgb = RGBColor(100, 116, 139)
        timeline_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    if title:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(15, 23, 42)
        _set_east_asia_font(run)
        paragraph.paragraph_format.space_after = Pt(10)

    if include_timeline:
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.style = "Table Grid"
        table.columns[0].width = Cm(3.5)
        table.columns[1].width = Cm(13.5)
        header = table.rows[0].cells
        header[0].text = "시간"
        header[1].text = "내용"
        for cell in header:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            # Soft blue-gray header shading; no theme dependency.
            tc_pr = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EFF6FF")
            tc_pr.append(shading)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 64, 175)
                _set_east_asia_font(run)

        for item in normalized:
            cells = table.add_row().cells
            stamp = f"{_format_time(item['start'])} – {_format_time(item['end'])}"
            timeline_p = cells[0].paragraphs[0]
            timeline_p.style = doc.styles["Timeline"]
            timeline_p.add_run(stamp)
            text_p = cells[1].paragraphs[0]
            text_p.paragraph_format.space_after = Pt(2)
            run = text_p.add_run(item["text"])
            _set_east_asia_font(run)
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    else:
        for item in normalized:
            paragraph = doc.add_paragraph(item["text"])
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                _set_east_asia_font(run)

    atomic_write(target, lambda temp: doc.save(str(temp)))
    return str(target)


# ---------------------------------------------------------------------------
# Refreshed CustomTkinter controls
# ---------------------------------------------------------------------------
class ModernButton(ctk.CTkButton):
    def __init__(self, master, **kwargs):
        super().__init__(
            master,
            corner_radius=12,
            height=44,
            font=ctk.CTkFont(size=14, weight="bold"),
            fg_color=PALETTE["primary"],
            hover_color=PALETTE["primary_hover"],
            text_color="#FFFFFF",
            border_width=0,
            **kwargs,
        )


class SecondaryButton(ctk.CTkButton):
    def __init__(self, master, **kwargs):
        super().__init__(
            master,
            corner_radius=12,
            height=42,
            font=ctk.CTkFont(size=13, weight="bold"),
            fg_color=PALETTE["primary_soft"],
            hover_color="#DBEAFE",
            text_color=PALETTE["primary_hover"],
            border_width=1,
            border_color="#BFDBFE",
            **kwargs,
        )


core.ModernButton = ModernButton
core.SecondaryButton = SecondaryButton


class App(core.App):
    """Polished shell around the existing feature panels."""

    def __init__(self):
        super().__init__()
        self.title(f"Study Helper {APP_VERSION}")

    def _build_ui(self):
        self.configure(fg_color=PALETTE["page"])
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.sidebar = ctk.CTkFrame(self, width=260, corner_radius=0, fg_color=PALETTE["sidebar"])
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_propagate(False)
        self.sidebar.grid_rowconfigure(2, weight=1)

        brand = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        brand.grid(row=0, column=0, padx=22, pady=(28, 18), sticky="ew")
        badge = ctk.CTkLabel(
            brand, text="SH", width=44, height=44, corner_radius=12,
            fg_color=PALETTE["primary"], text_color="#FFFFFF",
            font=ctk.CTkFont(size=18, weight="bold")
        )
        badge.grid(row=0, column=0, rowspan=2, padx=(0, 12))
        ctk.CTkLabel(
            brand, text="Study Helper", text_color="#FFFFFF",
            font=ctk.CTkFont(size=22, weight="bold")
        ).grid(row=0, column=1, sticky="sw")
        ctk.CTkLabel(
            brand, text=f"Audio study suite · v{APP_VERSION}",
            text_color=PALETTE["sidebar_muted"], font=ctk.CTkFont(size=11)
        ).grid(row=1, column=1, sticky="nw")

        ctk.CTkLabel(
            self.sidebar, text="WORKFLOWS", text_color=PALETTE["sidebar_muted"],
            font=ctk.CTkFont(size=11, weight="bold")
        ).grid(row=1, column=0, padx=22, pady=(8, 8), sticky="w")

        self.nav_scroll = ctk.CTkScrollableFrame(
            self.sidebar, fg_color="transparent", scrollbar_button_color="#334155",
            scrollbar_button_hover_color="#475569"
        )
        self.nav_scroll.grid(row=2, column=0, sticky="nsew", padx=12, pady=(0, 14))

        self.nav_buttons = {}
        self._nav_btn("YouTube → MP3", "youtube")
        self._nav_btn("MP3 → Excel · 번역/분리", "excel_trans")
        self._nav_btn("MP3 → Excel · 단순 STT", "excel_raw")
        self._nav_btn("MP3 → DOCX · 전사문", "pdf")
        self._nav_btn("Excel → MP3 · 한/영 교차", "tts_bilingual")
        self._nav_btn("Excel → MP3 · 3회 반복", "tts_3times")
        self._nav_btn("Excel → MP3 · 단일 음성", "tts_single")
        self._nav_btn("Excel → Anki", "anki")

        footer = ctk.CTkFrame(self.sidebar, fg_color="#111C30", corner_radius=14)
        footer.grid(row=3, column=0, padx=16, pady=(0, 18), sticky="ew")
        ctk.CTkLabel(
            footer, text="Local-first\nWhisper · Edge TTS · Anki",
            justify="left", text_color=PALETTE["sidebar_muted"],
            font=ctk.CTkFont(size=11)
        ).pack(anchor="w", padx=14, pady=12)

        self.main = ctk.CTkFrame(self, corner_radius=0, fg_color=PALETTE["page"])
        self.main.grid(row=0, column=1, sticky="nsew")
        self.main.grid_rowconfigure(1, weight=1)
        self.main.grid_columnconfigure(0, weight=1)

        self.header = ctk.CTkFrame(
            self.main, height=96, corner_radius=18, fg_color=PALETTE["card"],
            border_width=1, border_color=PALETTE["border"]
        )
        self.header.grid(row=0, column=0, sticky="ew", padx=24, pady=(22, 12))
        self.header.grid_columnconfigure(0, weight=1)
        self.title_label = ctk.CTkLabel(
            self.header, text="", font=ctk.CTkFont(size=26, weight="bold"),
            text_color=PALETTE["text"]
        )
        self.title_label.grid(row=0, column=0, padx=24, pady=(18, 2), sticky="w")
        self.desc_label = ctk.CTkLabel(
            self.header, text="", font=ctk.CTkFont(size=13), text_color=PALETTE["muted"]
        )
        self.desc_label.grid(row=1, column=0, padx=24, pady=(0, 18), sticky="w")
        self.status_badge = ctk.CTkLabel(
            self.header, text="준비", corner_radius=999, fg_color=PALETTE["success_bg"],
            text_color=PALETTE["success"], font=ctk.CTkFont(size=12, weight="bold"),
            padx=16, pady=8
        )
        self.status_badge.grid(row=0, column=1, rowspan=2, padx=22, pady=18, sticky="e")

        self.body = ctk.CTkScrollableFrame(
            self.main, fg_color="transparent", scrollbar_button_color="#CBD5E1",
            scrollbar_button_hover_color="#94A3B8"
        )
        self.body.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0, 22))
        self.body.grid_columnconfigure(0, weight=1)

        self.paste_widgets = []
        self.panel_builders = {
            "youtube": self._youtube_panel,
            "excel_trans": self._excel_panel_trans,
            "excel_raw": self._excel_panel_raw,
            "pdf": self._docx_panel,
            "tts_bilingual": self._tts_panel_bilingual,
            "tts_3times": self._tts_panel_3times,
            "tts_single": self._tts_panel_single,
            "anki": self._anki_panel,
        }
        self.panels = {}

    def _nav_btn(self, text, panel):
        btn = ctk.CTkButton(
            self.nav_scroll, text=text, command=lambda p=panel: self.show_panel(p),
            corner_radius=10, height=42, anchor="w", fg_color="transparent",
            hover_color=PALETTE["sidebar_hover"], text_color=PALETTE["sidebar_text"],
            font=ctk.CTkFont(size=13, weight="bold"), border_width=0
        )
        btn.pack(fill="x", pady=2)
        self.nav_buttons[panel] = btn

    def _panel_card(self, title, desc):
        frame = ctk.CTkFrame(
            self.body, corner_radius=18, fg_color=PALETTE["card"],
            border_width=1, border_color=PALETTE["border"]
        )
        top = ctk.CTkFrame(frame, fg_color="transparent")
        top.pack(fill="x", padx=28, pady=(24, 18))
        ctk.CTkLabel(
            top, text=title, font=ctk.CTkFont(size=21, weight="bold"),
            text_color=PALETTE["text"]
        ).pack(anchor="w")
        ctk.CTkLabel(
            top, text=desc, justify="left", wraplength=1180,
            text_color=PALETTE["muted"], font=ctk.CTkFont(size=13)
        ).pack(anchor="w", pady=(6, 0))
        return frame

    def section_label(self, parent, text):
        return ctk.CTkLabel(
            parent, text=text, font=ctk.CTkFont(size=13, weight="bold"),
            text_color=PALETTE["text"]
        )

    def input_box(self, parent):
        return ctk.CTkEntry(
            parent, height=44, corner_radius=10, border_width=1,
            border_color="#CBD5E1", fg_color=PALETTE["input"], text_color=PALETTE["text"]
        )

    def text_box(self, parent, height=220):
        return ctk.CTkTextbox(
            parent, height=height, corner_radius=12, border_width=1,
            border_color="#CBD5E1", fg_color=PALETTE["input"], text_color=PALETTE["text"]
        )

    def set_status(self, text, mode="normal"):
        def _update():
            self.status_badge.configure(text=text)
            if mode == "error":
                self.status_badge.configure(fg_color=PALETTE["error_bg"], text_color=PALETTE["error"])
            elif mode == "success":
                self.status_badge.configure(fg_color=PALETTE["success_bg"], text_color=PALETTE["success"])
            else:
                self.status_badge.configure(fg_color=PALETTE["primary_soft"], text_color=PALETTE["primary_hover"])
        self.post_ui(_update)

    def show_panel(self, name):
        titles = {
            "youtube": ("YouTube 링크를 MP3로 저장", "영상에서 음성만 추출해 학습용 MP3로 저장합니다."),
            "excel_trans": ("MP3를 Excel 카드로 변환", "문장 분리 · 한/영 번역 · 중복 제거를 한 번에 처리합니다."),
            "excel_raw": ("MP3를 Excel로 변환", "번역 없이 음성을 문장 단위로 전사해 Excel에 저장합니다."),
            "pdf": ("MP3를 DOCX 전사문으로 변환", "선택형 타임라인을 포함한 편집 가능한 Word 문서를 만듭니다."),
            "tts_bilingual": ("Excel을 MP3로 저장", "한/영 문장을 교차 재생하는 학습용 오디오를 생성합니다."),
            "tts_3times": ("Excel을 MP3로 저장 · 3회 반복", "미국·영국·호주 발음을 연속 재생하도록 오디오를 만듭니다."),
            "tts_single": ("Excel을 MP3로 저장 · 단일 음성", "선택한 국가와 성별의 음성으로 문장을 변환합니다."),
            "anki": ("Excel을 Anki Deck에 추가", "중복을 건너뛰고 새 학습 카드를 안전하게 추가합니다."),
        }
        title, desc = titles[name]
        self.title_label.configure(text=title)
        self.desc_label.configure(text=desc)
        if name not in self.panels:
            self.panels[name] = self.panel_builders[name]()
        for key, button in self.nav_buttons.items():
            if key == name:
                button.configure(fg_color=PALETTE["primary"], hover_color=PALETTE["primary_hover"], text_color="#FFFFFF")
            else:
                button.configure(fg_color="transparent", hover_color=PALETTE["sidebar_hover"], text_color=PALETTE["sidebar_text"])
        for panel in self.panels.values():
            panel.grid_forget()
        self.panels[name].grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=(0, 8))

    def _docx_panel(self):
        frame = self._panel_card(
            "MP3 → DOCX (Word 전사문)",
            "Whisper 전사 결과를 편집 가능한 DOCX로 저장합니다. 타임라인을 포함하면 왼쪽 첫 번째 열에만 배치되어 Word에서 그 열을 삭제하는 것만으로 타임라인 전체를 제거할 수 있습니다."
        )
        form = ctk.CTkFrame(frame, fg_color="transparent")
        form.pack(fill="both", expand=True, padx=28, pady=(0, 28))
        form.grid_columnconfigure(0, weight=1)

        self.section_label(form, "음성 파일").grid(row=0, column=0, sticky="w", pady=(0, 8))
        self.docx_audio = self.input_box(form)
        self.docx_audio.grid(row=1, column=0, sticky="ew")
        SecondaryButton(
            form, text="파일 찾기", width=104,
            command=lambda: self.pick_file_for(self.docx_audio, [("Audio", "*.mp3 *.wav *.m4a")])
        ).grid(row=1, column=1, padx=(12, 0), sticky="w")

        self.section_label(form, "저장 폴더").grid(row=2, column=0, sticky="w", pady=(18, 8))
        self.docx_dir = self.input_box(form)
        self.docx_dir.insert(0, str(Path.home() / "Documents"))
        self.docx_dir.grid(row=3, column=0, sticky="ew")
        SecondaryButton(
            form, text="폴더 변경", width=104,
            command=lambda: self.pick_dir_for(self.docx_dir)
        ).grid(row=3, column=1, padx=(12, 0), sticky="w")

        self.section_label(form, "저장 파일명 (미입력 시 원본 파일명)").grid(row=4, column=0, sticky="w", pady=(18, 8))
        self.docx_fn = self.input_box(form)
        self.docx_fn.grid(row=5, column=0, sticky="ew")

        option_box = ctk.CTkFrame(
            form, corner_radius=12, fg_color="#F8FAFC", border_width=1, border_color=PALETTE["border"]
        )
        option_box.grid(row=6, column=0, columnspan=2, sticky="ew", pady=(18, 0))
        self.docx_timeline_var = ctk.BooleanVar(value=True)
        self.docx_plain_copy_var = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(
            option_box, text="타임라인 포함 (권장)", variable=self.docx_timeline_var,
            fg_color=PALETTE["primary"], hover_color=PALETTE["primary_hover"],
            text_color=PALETTE["text"], font=ctk.CTkFont(size=13, weight="bold")
        ).pack(anchor="w", padx=16, pady=(14, 6))
        ctk.CTkCheckBox(
            option_box, text="타임라인 없는 사본도 함께 저장", variable=self.docx_plain_copy_var,
            fg_color=PALETTE["primary"], hover_color=PALETTE["primary_hover"],
            text_color=PALETTE["text"], font=ctk.CTkFont(size=13)
        ).pack(anchor="w", padx=16, pady=(0, 8))
        ctk.CTkLabel(
            option_box,
            text="타임라인 포함 문서는 1열=시간, 2열=내용 구조입니다. Word에서 첫 번째 열 전체 선택 → 열 삭제로 타임라인만 한 번에 제거할 수 있습니다.",
            justify="left", wraplength=1000, text_color=PALETTE["muted"], font=ctk.CTkFont(size=11)
        ).pack(anchor="w", padx=16, pady=(0, 14))

        ModernButton(
            form, text="DOCX 문서로 저장", width=206, command=self.run_stt_docx
        ).grid(row=7, column=0, columnspan=2, sticky="w", pady=(22, 0))
        return frame

    def run_stt_docx(self):
        audio_path = self.docx_audio.get().strip()
        directory = self.docx_dir.get().strip()
        filename = self.docx_fn.get().strip()
        include_timeline = bool(self.docx_timeline_var.get())
        plain_copy = bool(self.docx_plain_copy_var.get())

        def task():
            try:
                if not audio_path:
                    raise RuntimeError("음성 파일을 선택해 주세요.")
                if not directory:
                    raise RuntimeError("저장 폴더를 선택해 주세요.")
                self.set_status("음성 인식 중...")
                result = transcribe_audio(audio_path)
                target = Path(directory) / self.get_filename(audio_path, filename, ".docx")
                self.set_status("DOCX 생성 중...")
                speech_docx(result["segments"], target, include_timeline, Path(audio_path).stem)
                saved = [str(target)]
                if plain_copy and include_timeline:
                    plain_target = target.with_name(target.stem + "_no-timeline.docx")
                    speech_docx(result["segments"], plain_target, False, Path(audio_path).stem)
                    saved.append(str(plain_target))
                self.set_status("완료", "success")
                self.post_ui(lambda: messagebox.showinfo("완료", "저장 완료:\n" + "\n".join(saved)))
            except Exception as exc:
                logging.exception("DOCX 작업 실패")
                self.set_status("오류", "error")
                self.post_ui(lambda exc=exc: messagebox.showerror("오류", str(exc)))

        self.run_in_thread(task)


# Old key name remains for compatibility with saved navigation/state and tests.
App._pdf_panel = App._docx_panel


def _run_self_test(destination: str) -> int:
    from smoke_test import run
    try:
        run(sys.modules[__name__], destination)
    except Exception:
        return 1
    return 0


if __name__ == "__main__":
    multiprocessing.freeze_support()
    if len(sys.argv) == 3 and sys.argv[1] == "--self-test":
        code = _run_self_test(sys.argv[2])
        if getattr(sys, "frozen", False):
            os._exit(code)
        raise SystemExit(code)
    app = App()
    app.mainloop()
